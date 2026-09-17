"""Add-only import of the three approved supplier workbooks.

Run ``python -m scripts.catalog_import FILE ...`` for an offline preview. Supply
``--database-env DATABASE_URL`` (or ``--database-url``) for an exact database diff;
only ``--apply`` commits it. Environment access is always explicitly opted into.
No settings/.env are loaded, no seed functions run, and no offers are created.
Prices with an unknown currency or conflicting pack basis stay in audit metadata.
Ranges in fastener lists are held for review, never expanded into invented SKUs.
"""

from __future__ import annotations

import argparse
import asyncio
import hashlib
import json
import os
import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import insert, select, text
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from app.db.models.catalog import CanonicalProduct, Category, Unit

# Stable signed bigint namespace shared by every invocation of this importer.
IMPORT_LOCK = int.from_bytes(b"qurbotxl", "big")
LAYOUTS = {"fanera.xlsx": "sheet", "reka.xlsx": "timber", "samarez mix.xlsx": "fastener"}
UNITS = {
    "шт": "dona",
    "dona": "dona",
    "кг": "kg",
    "коробкa": "quti",
    "коробка": "quti",
    "пачка": "pachka",
}
CATEGORIES = {"sheet": "plita-va-fanera", "timber": "yogoch", "fastener": "mahkamlash-materiallari"}

# Literal workbook titles -> transcribed scripts.seed._METIZ identities.
# Deliberately explicit: neither fuzzy name matching nor price/currency evidence.
# Unit is part of the source identity (the white roofing screws have two groups).
FASTENER_SOURCES: dict[tuple[str, str], tuple[str, str | None]] = {
    ("ОК АНКЕР", "dona"): ("oq-anker", None),
    ("САРИК АНКЕР", "dona"): ("sariq-anker", None),
    ("АНКЕР КЛИН", "dona"): ("anker-klin", None),
    ("МУФТА СOEДИНИТЕЛ", "dona"): ("mufta-soedinitel", None),
    ("КРОВЕЛЬНЫЙ САМОРЕЗ РАНГЛИ", "quti"): ("krovelniy-samorez-rangli", None),
    ("КРОВЕЛЬНЫЙ САМОРЕЗ ОК", "quti"): ("krovelniy-samorez-oq", None),
    ("КРОВЕЛЬНЫЙ САМОРЕЗ ОК", "kg"): ("krovelniy-samorez-oq-kg", None),
    ("ДЮБЕЛ ГВОЗД", "kg"): ("dyubel-gvozd", None),
    ("МИХ", "kg"): ("mix", None),
    ("ЗАКЛЕПКА ORBITA", "pachka"): ("zaklepka-orbita", "Orbita"),
    ("РЕЗИНА ШАЙБА", "kg"): ("rezina-shayba", None),
    ("ПОДВЕС / АГРАФ", "quti"): ("podves-agraf", None),
    ("ЕВРО СТАШКА", "quti"): ("evro-stashka", None),
    ("ЦАНГА ORBITA", "dona"): ("tsanga-orbita", "Orbita"),
    ("ЦАНГА ТУРЕЦКИЙ", "dona"): ("tsanga-turk", None),
    ("КРЮЧОК САРИК", "dona"): ("kryuchok-sariq", None),
    ("КРЮЧОК ЁПИК", "dona"): ("kryuchok-yopiq", None),
    ("КРЮЧОК ОК ЁГОЧГА", "dona"): ("kryuchok-oq-yogochga", None),
    ("ГАЙКА ОК", "kg"): ("gayka-oq", None),
    ("ШПИЛЬКА 0,97", "dona"): ("shpilka-097", None),
    ("ЧОПИК СЕРЫЙ М / ПЛАСТ", "quti"): ("chopiq-kulrang", None),
    ("ЧОПИК КИЗИЛ М / ПЛАСТ", "quti"): ("chopiq-qizil", None),
    ("ЗАБИВНОЙ АНКЕР", "dona"): ("zabivnoy-anker", None),
    ("АНКЕР КРЮЧОК", "dona"): ("anker-kryuchok", None),
    ("АНКЕР КРЮЧОК ЁПИК", "dona"): ("anker-kryuchok-yopiq", None),
    ("ШПИЛЬКА 1метр ОРИГИНАЛ", "dona"): ("shpilka-1m", None),
    ("ШПИЛЬКА 2 метр", "dona"): ("shpilka-2m", None),
    ("ЗОНТИК М / ПЛАСТ", "pachka"): ("zontik-mplast", None),
    ("РЕЗБА ЗАКЛЕПКА", "dona"): ("rezba-zaklepka", None),
    ("КРЮЧОК АРМСТРОН", "dona"): ("kryuchok-armstrong", "Armstrong"),
    ("ГАЗАБЛОК ПРОПКА", "dona"): ("gazoblok-probka", None),
    ("КРЮЧОК БАБОЧКА", "dona"): ("kryuchok-babochka", None),
}
FASTENER_LABELS = {
    "бабочка": "babochka",
    "дрива": "driva",
    "зонтик шай": "zontik shay",
    "термо 120": "termo 120",
    "термо 150": "termo 150",
}


def norm(value: object) -> str:
    """Only typographic normalization; no fuzzy, transliteration or unit guesses."""
    value = "" if value is None else str(value)
    return " ".join(value.casefold().replace("×", "x").replace("х", "x").split())


def cell(value: object) -> str:
    return "" if value is None else str(value).strip()


@dataclass
class ImportRow:
    name: str
    category: str
    unit: str
    variant: dict[str, str]
    attributes: dict[str, Any]
    provenance: dict[str, Any]
    reference_price: Decimal | None = None
    brand: str | None = None
    legacy_slug: str | None = None
    legacy_attributes: dict[str, Any] = field(default_factory=dict)
    legacy_source_ref: str = "fanera.uz"

    @property
    def key(self) -> str:
        identity = [self.category, self.unit, norm(self.brand), self.variant]
        return hashlib.sha256(json.dumps(identity, sort_keys=True).encode()).hexdigest()

    @property
    def slug(self) -> str:
        # Recognized legacy identities must also be stable if import precedes
        # the deploy seed, otherwise that seed would recreate the same variant.
        return self.legacy_slug or f"excel-{self.key}"


def _price(value: str) -> Decimal | None:
    try:
        number = Decimal(value.replace(" ", "").replace(",", "."))
    except InvalidOperation:
        return None
    if not number.is_finite() or number <= 0 or number >= Decimal("1000000000000"):
        return None
    return number if number == number.quantize(Decimal("0.01")) else None


def _sheet_legacy(row: ImportRow, values: list[str]) -> None:
    """Recognize the existing explicitly transcribed fanera.uz variants only."""
    name, origin, grade, size, thickness, _ = values
    size, grade = norm(size), norm(grade)
    material = ""
    suffix = thickness.replace(".", "-")
    if name == "Fanera" and origin == "Rossiya":
        row.legacy_slug = f"fanera-bereza-{grade}-{suffix}mm-{size}"
        material = "birch_plywood"
        row.legacy_attributes["grade"] = grade
    elif name.startswith("Fanera laminatsiyalangan"):
        brand = name.removeprefix("Fanera laminatsiyalangan").strip()
        if origin == "Rossiya" and brand in {"SibPly", "SiyPly", "Murashinsky", "SEGEZHA"}:
            row.brand = brand
            row.legacy_slug = f"fanera-laminat-{brand.lower()}-{suffix}mm-{size}"
            material = "laminated_plywood"
        elif origin == "Xitoy" and not brand:
            row.legacy_slug = f"fanera-laminat-xitoy-{suffix}mm-{size}"
            material = "laminated_plywood"
    elif name == "OSB-3" and origin == "Rossiya":
        row.legacy_slug = f"osb3-{suffix}mm-{size}"
        material = "osb3"
    elif (name, origin, grade) in {
        ("HDF", "Rossiya (Kronospan)", ""),
        ("DVP (T markasi)", "Rossiya (Perm)", "-"),
    }:
        nominal = re.fullmatch(r"(\d+(?:\.\d+)?) \((\d+)\)", thickness)
        if nominal:
            thickness, trade = nominal.groups()
            suffix = thickness.replace(".", "-")
            row.legacy_attributes["thickness_trade_mm"] = trade
            if name == "HDF":
                row.brand = "Kronospan"
                material = "hdf"
                row.legacy_slug = f"hdf-kronospan-{suffix}mm-{size}"
            else:
                material = "hardboard"
                row.legacy_slug = f"dvp-t-{suffix}mm-{size}"
    if material:
        row.legacy_attributes.update(material=material, size=size, thickness_mm=thickness)


def _fastener_legacy(row: ImportRow, name: str, size: str) -> str | None:
    """Return a hold reason unless the workbook family and size are understood."""
    source = next(
        (
            value
            for (title, unit), value in FASTENER_SOURCES.items()
            if norm(title) == norm(name) and unit == row.unit
        ),
        None,
    )
    if source is None:
        return "unknown/conflicting fastener source group or unit; confirm transcription"
    group, row.brand = source
    label = re.sub(r"(?<=\d),(?=\d)", ".", norm(size))
    label = re.sub(r"^[mм]\s*(?=\d)", "m", label)
    label = FASTENER_LABELS.get(label, label)
    if (
        not re.fullmatch(r"m?\d+(?:\.\d+)?(?:x\d+(?:\.\d+)?){0,2}", label)
        and label not in FASTENER_LABELS.values()
    ):
        return "unrecognized fastener size label; confirm transcription"
    row.legacy_source_ref = "metiz-prays"
    row.legacy_slug = group + "-" + re.sub(r"[^a-z0-9]+", "-", label).strip("-")
    row.legacy_attributes = {"material": group.replace("-", "_")}
    if "x" in label:
        row.legacy_attributes["size"] = label
    head = re.match(r"^m?(\d+(?:\.\d+)?)", label)
    if head:
        row.legacy_attributes["diameter_mm"] = head.group(1)
    if re.fullmatch(r"m\d+", label):
        row.legacy_attributes["grade"] = label
    row.attributes.update(row.legacy_attributes)
    row.provenance["source_group"] = group
    row.provenance["source_size"] = label
    return None


def read_workbook(path: Path) -> tuple[list[ImportRow], list[dict[str, Any]]]:
    """Parse a single immutable byte snapshot; checksum refers to those exact bytes."""
    layout = LAYOUTS.get(path.name)
    if layout is None:
        raise ValueError(f"Unsupported workbook: {path.name}")
    payload = path.read_bytes()
    checksum = hashlib.sha256(payload).hexdigest()
    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=False)
    rows: list[ImportRow] = []
    review: list[dict[str, Any]] = []
    try:
        for sheet in workbook:
            header_seen = False
            for number, raw in enumerate(sheet.iter_rows(values_only=True), 1):
                values = [cell(v) for v in raw]
                if not any(values):
                    continue
                values += [""] * max(0, 6 - len(values))
                provenance = {
                    "file": path.name,
                    "sha256": checksum,
                    "sheet": sheet.title,
                    "row": number,
                    "cells": values,
                }
                expected = "Наименования" if layout == "fastener" else "Mahsulot nomi"
                if expected in values:
                    header_seen = True
                    # Validate the price basis too, rather than trusting a filename.
                    if layout != "fastener" and values[5] not in {
                        "Narx (So’m/dona)",
                        "Narx (So'm/dona)",
                    }:
                        raise ValueError(f"Unexpected price header: {path.name}:{sheet.title}")
                    continue
                if not header_seen:
                    review.append({**provenance, "reason": "unrecognized header or preamble"})
                    continue
                if any(v.startswith("=") for v in values):
                    review.append({**provenance, "reason": "formula requires review"})
                    continue
                attrs: dict[str, Any] = {}
                if layout == "sheet":
                    name, origin, grade, size, thickness, price = values[:6]
                    # These exact source discrepancies are already held in seed.py.
                    # Never repair 1.6 -> 16 or 2071 -> 2070 by inference.
                    if (norm(name) == "dsp" and norm(thickness) == "1.6") or (
                        norm(name) == "hdf" and norm(grade) == "oq" and norm(size) == "2800x2071"
                    ):
                        review.append(
                            {
                                **provenance,
                                "reason": "supplier-held sheet dimensions require confirmation",
                            }
                        )
                        continue
                    unit = "dona"
                    variant = dict(
                        name=norm(name),
                        origin=norm(origin),
                        grade=norm(grade),
                        size=norm(size),
                        thickness_mm=norm(thickness),
                    )
                    attrs.update(
                        origin=origin, grade=grade, size=norm(size), thickness_label=thickness
                    )
                    if re.fullmatch(r"\d+(?:\.\d+)?", thickness):
                        attrs["thickness_mm"] = thickness
                    label = " ".join(v for v in [name, origin, grade, size, f"{thickness} mm"] if v)
                    reference_price = _price(price)
                elif layout == "timber":
                    name, size, pack, length, grade, price = values[:6]
                    unit = UNITS.get(norm(pack), "")
                    variant = dict(
                        name=norm(name),
                        size=norm(size),
                        pack=norm(pack),
                        length=norm(length),
                        grade=norm(grade),
                    )
                    attrs.update(size=norm(size), grade=grade, pack_label=pack, length_label=length)
                    label = " ".join([name, size, pack, length, grade])
                    # The printed dona header conflicts with пачка on bundle rows.
                    reference_price = _price(price) if unit == "dona" else None
                else:
                    _, name, size, pack, price = values[:5]
                    unit = UNITS.get(norm(pack), "")
                    variant = dict(name=norm(name), size=norm(size))
                    attrs["size"] = norm(size)
                    label = f"{name} {size}"
                    reference_price = None  # No currency is printed in this workbook.
                    if ";" in size or re.search(r"\d\s*[-–/]", size):
                        review.append({**provenance, "reason": "size range/list requires review"})
                        continue
                if (
                    not name
                    or not size
                    or not unit
                    or len(label) > 255
                    or (layout == "sheet" and (not origin or not thickness))
                ):
                    review.append({**provenance, "reason": "missing/unsupported identity or unit"})
                    continue
                attrs.update(price_on_request=reference_price is None, stock_unverified=True)
                provenance.update(
                    raw_price=price,
                    currency="UZS" if layout != "fastener" else None,
                    price_basis_verified=layout == "sheet" or unit == "dona",
                )
                row = ImportRow(
                    label, CATEGORIES[layout], unit, variant, attrs, provenance, reference_price
                )
                if layout == "sheet":
                    _sheet_legacy(row, values[:6])
                elif layout == "fastener":
                    reason = _fastener_legacy(row, name, size)
                    if reason:
                        review.append({**provenance, "reason": reason})
                        continue
                rows.append(row)
    finally:
        workbook.close()
    return rows, review


def _exact(row: ImportRow, product: dict[str, Any]) -> bool:
    attrs = product["attributes"] or {}
    if product["base_unit_code"] != row.unit or norm(product["brand"]) != norm(row.brand):
        return False
    import_meta = attrs.get("catalog_import")
    if (
        isinstance(import_meta, dict)
        and import_meta.get("variant") == row.variant
        and all(
            norm(attrs.get(k)) == norm(v)
            for k, v in row.attributes.items()
            if k not in {"price_on_request", "stock_unverified"}
        )
    ):
        return True
    if (
        row.legacy_slug == product["slug"]
        and product["source_ref"] == row.legacy_source_ref
        and all(norm(attrs.get(k)) == norm(v) for k, v in row.legacy_attributes.items())
    ):
        return True
    # Existing manually entered variants require every supplied specification.
    return any(
        norm(product[k]) == norm(row.name) for k in ("name_uz", "name_uz_cyrl", "name_ru")
    ) and all(
        norm(attrs.get(k)) == norm(v)
        for k, v in row.attributes.items()
        if k not in {"price_on_request", "stock_unverified"}
    )


def _source_collision(row: ImportRow, product: dict[str, Any]) -> bool:
    """Related source identity is a review signal, never evidence to auto-merge."""
    if product["slug"] == row.slug:
        return True
    attrs = product["attributes"] or {}
    expected = row.legacy_attributes
    if not expected or product["source_ref"] != row.legacy_source_ref:
        return False
    if norm(attrs.get("material")) != norm(expected.get("material")):
        return False
    # Detect renamed legacy identities with conflicting brand, unit or details.
    # Full size (not diameter alone) is required whenever the source supplies it.
    dimensions = ("size", "thickness_mm", "thickness_trade_mm", "grade")
    relevant = {key: expected[key] for key in dimensions if key in expected}
    if not relevant and "diameter_mm" in expected:
        relevant["diameter_mm"] = expected["diameter_mm"]
    return bool(relevant) and all(norm(attrs.get(k)) == norm(v) for k, v in relevant.items())


async def import_rows(
    session: AsyncSession, rows: list[ImportRow], *, apply: bool = False
) -> list[dict[str, Any]]:
    """Own one transaction; dry runs execute SELECT only. Existing rows are immutable."""
    if session.in_transaction():
        raise ValueError("Importer requires a fresh session/transaction")
    decisions: list[dict[str, Any]] = []
    async with session.begin():
        dialect = session.get_bind().dialect.name
        if apply:
            if dialect == "postgresql":
                await session.execute(
                    text("SELECT pg_advisory_xact_lock(:key)"), {"key": IMPORT_LOCK}
                )
            elif dialect == "sqlite":
                await session.execute(text("BEGIN IMMEDIATE"))
            else:
                raise ValueError("Apply supports PostgreSQL and SQLite only")
        elif dialect == "postgresql":
            await session.execute(text("SET TRANSACTION READ ONLY"))
        categories = dict((await session.execute(select(Category.slug, Category.id))).all())
        units = set((await session.execute(select(Unit.code))).scalars())
        products = [
            dict(p) for p in (await session.execute(select(CanonicalProduct.__table__))).mappings()
        ]
        planned: dict[str, int | None] = {}
        sources: dict[str, list[dict[str, Any]]] = {}
        for row in rows:
            sources.setdefault(row.key, []).append(row.provenance)
        for row in rows:
            decision = {
                **row.provenance,
                "key": row.key,
                "name": row.name,
                "reference_price": str(row.reference_price) if row.reference_price else None,
            }
            matches = [
                p
                for p in products
                if p["category_id"] == categories.get(row.category) and _exact(row, p)
            ]
            collisions = [p["id"] for p in products if _source_collision(row, p)]
            if len(matches) > 1:
                decision.update(action="review", reason="multiple exact existing variants")
            elif matches:
                decision.update(action="existing", canonical_id=matches[0]["id"])
            elif row.key in planned:
                decision.update(action="duplicate", canonical_id=planned[row.key])
            elif row.category not in categories or row.unit not in units:
                decision.update(action="review", reason="required category or unit is missing")
            elif collisions:
                decision.update(
                    action="review",
                    reason="source identity conflicts with existing details",
                    candidate_ids=collisions,
                )
            else:
                decision.update(action="add")
                planned[row.key] = None
                if apply:
                    attrs = {
                        **row.attributes,
                        "catalog_import": {
                            "version": 1,
                            "variant": row.variant,
                            "sources": sources[row.key],
                        },
                    }
                    result = await session.execute(
                        insert(CanonicalProduct)
                        .values(
                            slug=row.slug,
                            name_uz=row.name,
                            name_uz_cyrl=row.name,
                            name_ru=row.name,
                            brand=row.brand,
                            category_id=categories[row.category],
                            base_unit_code=row.unit,
                            attributes=attrs,
                            tier="standard",
                            source="supplier",
                            source_ref=f"sha256:{row.provenance['sha256']}",
                            reference_price=row.reference_price,
                            is_active=True,
                            search_doc=norm(row.name),
                        )
                        .returning(CanonicalProduct.id)
                    )
                    planned[row.key] = result.scalar_one()
                    decision["canonical_id"] = planned[row.key]
            decisions.append(decision)
    return decisions


async def run(args: argparse.Namespace) -> dict[str, Any]:
    database_url = args.database_url
    database_env = getattr(args, "database_env", None)
    if database_env:
        if database_url:
            raise ValueError("--database-env and --database-url are mutually exclusive")
        database_url = os.environ.get(database_env)
        if not database_url:
            raise ValueError("Explicit --database-env names an unset or empty environment variable")
    rows: list[ImportRow] = []
    review: list[dict[str, Any]] = []
    for path in args.files:
        parsed, held = await asyncio.to_thread(read_workbook, path)
        rows.extend(parsed)
        review.extend(held)
    if args.apply and not database_url:
        raise ValueError("--apply requires an explicit --database-url or --database-env")
    if database_url:
        engine = create_async_engine(database_url)
        try:
            async with async_sessionmaker(engine)() as session:
                decisions = await import_rows(session, rows, apply=args.apply)
        finally:
            await engine.dispose()
    else:
        decisions = [
            {
                **r.provenance,
                "key": r.key,
                "name": r.name,
                "action": "candidate",
                "reference_price": str(r.reference_price) if r.reference_price else None,
                "attributes": r.attributes,
            }
            for r in rows
        ]
    counts = {
        action: sum(d["action"] == action for d in decisions)
        for action in sorted({d["action"] for d in decisions})
    }
    return {
        "mode": "apply" if args.apply else "dry-run",
        "database_compared": bool(database_url),
        "counts": counts,
        "review_count": len(review),
        "review": review,
        "rows": decisions,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("files", type=Path, nargs="+")
    database = parser.add_mutually_exclusive_group()
    database.add_argument(
        "--database-url", help="Explicit SQLAlchemy async URL; never read from .env"
    )
    database.add_argument(
        "--database-env",
        metavar="VARIABLE",
        help="Explicitly opt in to a named database URL environment variable",
    )
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument("--apply", action="store_true", help="Commit new canonical products only")
    mode.add_argument("--dry-run", action="store_true", help="Preview only (the default)")
    args = parser.parse_args()
    print(json.dumps(asyncio.run(run(args)), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
